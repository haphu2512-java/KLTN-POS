"""Xuất kết quả: file Excel, biểu đồ PNG, báo cáo tiếng Việt cho quản lý."""
from __future__ import annotations

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import config as C  # noqa: E402

# Bảng màu (đã kiểm tra phân biệt được với người mù màu): xanh = thực tế, cam = dự báo.
INK, MUTED, GRID = "#0b0b0b", "#8a8984", "#e4e3df"
ACTUAL, FORECAST = "#2a78d6", "#eb6834"
plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 9, "axes.edgecolor": MUTED, "axes.labelcolor": INK,
    "xtick.color": "#52514e", "ytick.color": "#52514e", "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6, "axes.axisbelow": True,
})


def vnd(x: float) -> str:
    """Định dạng tiền: 12.3 triệu đồng / 1,25 tỷ đồng."""
    if abs(x) >= 1e9:
        return f"{x / 1e9:,.2f} tỷ đồng".replace(",", "X").replace(".", ",").replace("X", ".")
    if abs(x) >= 1e6:
        return f"{x / 1e6:,.1f} triệu đồng".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{x / 1e3:,.0f} nghìn đồng".replace(",", ".")


def num(x: float) -> str:
    return f"{x:,.0f}".replace(",", ".")


def pct(x: float) -> str:
    return f"{x * 100:.1f}%".replace(".", ",")


# ================================ BIỂU ĐỒ ==========================================
def _tet_spans(ax, start, end):
    for tet in C.TET_DATES:
        a, b = tet - pd.Timedelta(days=21), tet + pd.Timedelta(days=9)
        if b >= start and a <= end:
            ax.axvspan(max(a, start), min(b, end), color="#f4ecd9", zorder=0, lw=0)
            ax.text(tet, ax.get_ylim()[1] * 0.97, "Tết", ha="center", va="top", fontsize=8, color="#52514e")


def chart_overview(R: dict, path) -> None:
    w, X, fl = R["weekly"], R["X"], R["forecast_long"]
    tot = w.groupby("week_start")[["sales", "demand"]].sum()
    fut = fl.groupby("week_start")[["forecast"]].sum()
    # Khoảng tin cậy của tổng: cộng phương sai (giả định độc lập giữa các chuỗi).
    var = fl.assign(v=((fl["hi95"] - fl["forecast"]) / 1.96) ** 2).groupby("week_start")["v"].sum()
    fut["lo"], fut["hi"] = fut["forecast"] - 1.96 * np.sqrt(var), fut["forecast"] + 1.96 * np.sqrt(var)
    fig, ax = plt.subplots(figsize=(11, 4.2))
    ax.plot(tot.index, tot["sales"], color=MUTED, lw=1.2, ls="--", label="Số bán ghi nhận (bị hụt khi hết hàng)")
    ax.plot(tot.index, tot["demand"], color=ACTUAL, lw=2, label="Nhu cầu thực tế đã làm sạch")
    ax.fill_between(fut.index, fut["lo"], fut["hi"], color=FORECAST, alpha=0.18, lw=0, label="Khoảng tin cậy 95%")
    ax.plot(fut.index, fut["forecast"], color=FORECAST, lw=2, label="Dự báo 12 tuần")
    ax.set_ylim(0, max(tot["demand"].max(), fut["hi"].max()) * 1.1)
    _tet_spans(ax, tot.index.min(), fut.index.max())
    ax.set_title("Tổng nhu cầu theo tuần - 80 SKU, 2 chi nhánh (DỮ LIỆU MẪU)", loc="left", fontsize=11, color=INK)
    ax.set_ylabel("Đơn vị / tuần")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.09), frameon=False, ncol=4, fontsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def chart_models(R: dict, path) -> None:
    t = R["eval"]["overall"].sort_values("WAPE", ascending=False)
    fig, ax = plt.subplots(figsize=(7, 3.2))
    bars = ax.barh(t.index, t["WAPE"] * 100, color=ACTUAL, height=0.55)
    for b, v in zip(bars, t["WAPE"]):
        ax.text(b.get_width() + 0.5, b.get_y() + b.get_height() / 2, pct(v), va="center", fontsize=9, color=INK)
    ax.set_xlabel("WAPE (%) - càng thấp càng tốt")
    ax.set_title("Sai số dự báo 12 tuần trong backtest (3 lần cắt thời gian)", loc="left", fontsize=11, color=INK)
    ax.grid(axis="y", visible=False)
    ax.set_xlim(0, t["WAPE"].max() * 118)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _sku_panel(ax, sku: str, R: dict) -> None:
    w, fl, bt, sel = R["weekly"], R["forecast_long"], R["backtest"], R["series_final"]
    hist = w[w["sku"] == sku].groupby("week_start")["demand"].sum()
    hist = hist[hist.index >= hist.index.max() - pd.Timedelta(weeks=51)]
    f = fl[fl["sku"] == sku].groupby("week_start")[["forecast"]].sum()
    var = fl[fl["sku"] == sku].assign(v=((fl["hi95"] - fl["forecast"]) / 1.96) ** 2).groupby("week_start")["v"].sum()
    lo, hi = (f["forecast"] - 1.96 * np.sqrt(var)).clip(lower=0), f["forecast"] + 1.96 * np.sqrt(var)
    # Dự báo backtest gốc cuối (mô hình được chọn) để so với thực tế đã xảy ra.
    model = sel.loc[sel["sku"] == sku, "best_model"].iloc[0]
    last_o = C.BACKTEST_ORIGINS[-1]
    b = bt[(bt["sku"] == sku) & (bt["origin"] == last_o) & (bt["model"] == model)].groupby("h")["forecast"].sum()
    b_weeks = R["X"]["week_start"][last_o:last_o + len(b)]
    ax.plot(hist.index, hist.values, color=ACTUAL, lw=1.6, label="Thực tế")
    ax.plot(b_weeks, b.values, color=FORECAST, lw=1.4, ls="--", label="Dự báo backtest")
    ax.fill_between(f.index, lo, hi, color=FORECAST, alpha=0.18, lw=0, label="KTC 95%")
    ax.plot(f.index, f["forecast"], color=FORECAST, lw=1.8, label="Dự báo 12 tuần")
    ax.axvline(C.END_DATE, color=MUTED, lw=0.8)
    name = sel.loc[sel["sku"] == sku, "name"].iloc[0]
    ax.set_title(f"{sku} · {name[:38]}" + ("…" if len(name) > 38 else ""), loc="left", fontsize=8.5, color=INK)
    ax.text(0.99, 0.95, model, transform=ax.transAxes, ha="right", va="top", fontsize=7.5, color="#52514e")
    ax.set_ylim(bottom=0)
    ax.tick_params(axis="x", labelsize=7, rotation=0)
    ax.xaxis.set_major_formatter(matplotlib.dates.DateFormatter("%m/%y"))


def chart_top10(R: dict, top10: list[str]) -> list:
    paths = []
    fig, axes = plt.subplots(5, 2, figsize=(12, 14))
    for ax, sku in zip(axes.ravel(), top10):
        _sku_panel(ax, sku, R)
    h, l = axes[0, 0].get_legend_handles_labels()
    fig.suptitle("Top 10 SKU theo doanh thu: thực tế vs dự báo (tổng 2 chi nhánh, DỮ LIỆU MẪU)",
                 x=0.01, y=0.995, ha="left", fontsize=12, color=INK)
    fig.legend(h, l, loc="upper left", ncol=4, frameon=False, bbox_to_anchor=(0.005, 0.975))
    fig.tight_layout(rect=(0, 0, 1, 0.955))
    p = C.CHART_DIR / "top10_thuc_te_vs_du_bao.png"
    fig.savefig(p, dpi=130)
    plt.close(fig)
    paths.append(p)
    for i, sku in enumerate(top10, 1):
        fig, ax = plt.subplots(figsize=(9, 3.4))
        _sku_panel(ax, sku, R)
        ax.legend(loc="upper left", frameon=False, fontsize=8, ncol=4)
        fig.tight_layout()
        p = C.CHART_DIR / f"top{i:02d}_{sku}.png"
        fig.savefig(p, dpi=130)
        plt.close(fig)
        paths.append(p)
    return paths


# ================================ EXCEL ============================================
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
ALERT_FILL = {3: PatternFill("solid", fgColor="F8CBAD"), 2: PatternFill("solid", fgColor="FFE699"),
              1: PatternFill("solid", fgColor="DDEBF7")}


def _format_sheet(ws, df: pd.DataFrame, formats: dict | None = None) -> None:
    formats = formats or {}
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32
    for j, col in enumerate(df.columns, 1):
        letter = get_column_letter(j)
        sample = [str(col)] + [str(v) for v in df[col].head(200)]
        ws.column_dimensions[letter].width = min(max(len(s) for s in sample) + 2, 48)
        fmt = formats.get(col)
        if fmt is None and pd.api.types.is_float_dtype(df[col]):
            fmt = "#,##0.0"
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
    ws.auto_filter.ref = ws.dimensions


def write_excel(R: dict, path) -> None:
    inv, seg, ev, fl = R["inventory"], R["segments"], R["eval"], R["forecast_long"]
    store = {k: v for k, v in C.STORES.items()}
    sheets: list[tuple[str, pd.DataFrame, dict]] = []

    sheets.append(("Tóm tắt", pd.DataFrame(R["kpis"], columns=["Chỉ tiêu", "Giá trị"]), {}))
    sheets.append(("Giả định", pd.DataFrame(R["assumptions"], columns=["Hạng mục", "Giả định"]), {}))

    order = inv.sort_values(["alert_level", "order_value"], ascending=[False, False])
    o = pd.DataFrame({
        "Chi nhánh": order["store_id"].map(store), "Mã SKU": order["sku"], "Tên sản phẩm": order["name"],
        "Nhóm hàng": order["category"], "ABC-XYZ": order["segment"], "Mô hình": order["best_model"],
        "Tồn hiện có": order["on_hand"], "Hàng đang về": order["on_order"],
        "Dự báo tuần tới": order["f1"].round(1), "Dự báo TB/tuần (12 tuần)": order["avg_weekly_fc"].round(1),
        "Sai số dự báo σ (đv/tuần)": order["sigma_week"].round(1), "Tồn kho an toàn (SS)": order["safety_stock"],
        "Điểm đặt hàng lại (ROP)": order["ROP"], "Mức tồn mục tiêu (S)": order["order_up_to"],
        "Quy cách": order["pack_size"], "SL đề xuất đặt": order["order_qty"],
        "Giá vốn": order["cost_price"], "Giá trị đặt hàng": order["order_value"],
        "Số tuần tồn đủ bán": order["weeks_of_cover"].round(1), "Xác suất hết hàng trước khi hàng về": order["stockout_prob"],
        "Cảnh báo": order["alert"], "Hành động": order["action"], "_lvl": order["alert_level"],
    })
    sheets.append(("Đề xuất đặt hàng", o, {"Giá vốn": "#,##0", "Giá trị đặt hàng": "#,##0",
                                           "Xác suất hết hàng trước khi hàng về": "0.0%"}))

    al = o[o["_lvl"] > 0].copy()
    al["Vốn tồn dư (giá vốn)"] = inv.set_index(["store_id", "sku"]).loc[
        list(zip(order.loc[al.index, "store_id"], order.loc[al.index, "sku"])), "excess_value"].values
    al = al[["Chi nhánh", "Mã SKU", "Tên sản phẩm", "ABC-XYZ", "Cảnh báo", "Tồn hiện có", "Hàng đang về",
             "Dự báo tuần tới", "Số tuần tồn đủ bán", "Xác suất hết hàng trước khi hàng về", "SL đề xuất đặt",
             "Vốn tồn dư (giá vốn)", "Hành động", "_lvl"]]
    sheets.append(("Cảnh báo", al, {"Xác suất hết hàng trước khi hàng về": "0.0%", "Vốn tồn dư (giá vốn)": "#,##0"}))

    piv = fl.groupby(["sku", "name", "week_start"])["forecast"].sum().unstack().round(1)
    piv.columns = [f"T{i + 1} ({d:%d/%m})" for i, d in enumerate(piv.columns)]
    piv = piv.reset_index().rename(columns={"sku": "Mã SKU", "name": "Tên sản phẩm"})
    piv.insert(2, "Tổng 12 tuần", piv.iloc[:, 2:].sum(axis=1).round(1))
    sheets.append(("Dự báo 12 tuần (tổng)", piv, {}))

    f = fl.assign(store=fl["store_id"].map(store))[["store", "sku", "name", "segment", "model", "week",
                                                    "week_start", "forecast", "lo95", "hi95"]]
    f.columns = ["Chi nhánh", "Mã SKU", "Tên sản phẩm", "ABC-XYZ", "Mô hình", "Tuần thứ", "Từ ngày",
                 "Dự báo", "Cận dưới 95%", "Cận trên 95%"]
    sheets.append(("Dự báo chi tiết", f, {"Từ ngày": "dd/mm/yyyy"}))

    s = seg.rename(columns={"sku": "Mã SKU", "name": "Tên sản phẩm", "category": "Nhóm hàng",
                            "revenue_52w": "Doanh thu 52 tuần", "revenue_share": "Tỷ trọng DT",
                            "cum_share": "Tỷ trọng lũy kế", "weekly_mean": "Nhu cầu TB/tuần",
                            "CV": "Hệ số biến thiên (CV)", "CV2_size": "CV² cỡ đơn",
                            "demand_pattern": "Kiểu nhu cầu", "segment": "Nhóm ABC-XYZ"})
    s = s.merge(ev["selection"].rename(columns={"segment": "Nhóm ABC-XYZ", "best_model": "Mô hình được chọn"})[
        ["Nhóm ABC-XYZ", "Mô hình được chọn"]], on="Nhóm ABC-XYZ", how="left")
    sheets.append(("ABC-XYZ", s, {"Doanh thu 52 tuần": "#,##0", "Tỷ trọng DT": "0.0%", "Tỷ trọng lũy kế": "0.0%"}))

    def mtab(t: pd.DataFrame) -> pd.DataFrame:
        t = t.reset_index().rename(columns={"index": "Mô hình", "model": "Mô hình", "n_points": "Số điểm"})
        return t
    ov = mtab(ev["overall"])
    ho = mtab(ev["holdout"])
    ho["Mô hình"] = ho["Mô hình"].astype(str) + " (kiểm tra gốc cuối)"
    bs = ev["by_segment"].rename(columns={"segment": "Nhóm", "model": "Mô hình", "n_points": "Số điểm"})
    bs = bs.merge(ev["selection"][["segment", "best_model"]].rename(columns={"segment": "Nhóm"}), on="Nhóm")
    bs["Được chọn"] = np.where(bs["Mô hình"] == bs["best_model"], "✔", "")
    bs = bs.drop(columns="best_model")
    comp = pd.concat([ov.assign(Nhóm="Tất cả (3 gốc)"), ho.assign(Nhóm="Tất cả"), bs], ignore_index=True)
    comp = comp[["Nhóm", "Mô hình", "MAE", "WAPE", "Bias", "Số điểm", "Được chọn"]].fillna({"Được chọn": ""})
    sheets.append(("So sánh mô hình", comp, {"WAPE": "0.0%", "Bias": "+0.0%;-0.0%", "MAE": "#,##0.00",
                                             "Số điểm": "#,##0"}))

    q = R["quality"].assign(store=lambda d: d["store_id"].map(store))
    q = q[["store", "sku", "days", "missing_record", "missing_qty", "outliers", "stockout_days", "stockout_rate",
           "sales_raw", "demand_clean", "lost_sales_est"]]
    q.columns = ["Chi nhánh", "Mã SKU", "Số ngày mở cửa", "Mất bản ghi", "Thiếu số lượng", "Ngoại lai đã sửa",
                 "Ngày hết hàng", "Tỷ lệ ngày hết hàng", "Tổng bán ghi nhận", "Tổng nhu cầu đã làm sạch",
                 "Ước tính bán mất do hết hàng"]
    sheets.append(("Chất lượng dữ liệu", q, {"Tỷ lệ ngày hết hàng": "0.0%"}))

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df, fmts in sheets:
            lvl = df["_lvl"].values if "_lvl" in df else None
            out = df.drop(columns=["_lvl"]) if lvl is not None else df
            out.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            _format_sheet(ws, out, fmts)
            if lvl is not None:  # tô màu theo mức cảnh báo
                for r, lv in enumerate(lvl, start=2):
                    if lv in ALERT_FILL:
                        for c in ws[r]:
                            c.fill = ALERT_FILL[lv]
        xw.sheets["Tóm tắt"].column_dimensions["A"].width = 52
        xw.sheets["Tóm tắt"].column_dimensions["B"].width = 40
        xw.sheets["Giả định"].column_dimensions["B"].width = 110


# ================================ BÁO CÁO ==========================================
def write_report(R: dict, path) -> None:
    k, ev, inv, seg = R["facts"], R["eval"], R["inventory"], R["segments"]
    ov, ho = ev["overall"], ev["holdout"]
    best = ov["WAPE"].idxmin()
    base_wape = ov.loc["MA8", "WAPE"]
    urgent = inv[inv["alert_level"] >= 2].sort_values(["alert_level", "stockout_prob"], ascending=False)
    over = inv[inv["alert"].isin(["Tồn kho quá mức", "Hàng không luân chuyển"])].sort_values("excess_value", ascending=False)
    seg_counts = seg["segment"].value_counts()
    n_a = (seg["ABC"] == "A").sum()

    def row_u(r):
        return (f"| {C.STORES[r.store_id].split(' - ')[0]} | {r.sku} | {r.name[:45]} | {r.alert} | "
                f"{int(r.on_hand)} | {str(round(r.f1, 1)).replace('.', ',')} | {int(r.order_qty)} |")

    def row_o(r):
        return (f"| {C.STORES[r.store_id].split(' - ')[0]} | {r.sku} | {r.name[:45]} | {str(round(r.weeks_of_cover, 1)).replace('.', ',')} | "
                f"{vnd(r.excess_value)} |")

    lines = [
        "# Báo cáo dự báo nhu cầu và tồn kho",
        "",
        "> **Lưu ý: toàn bộ số liệu trong báo cáo này là DỮ LIỆU MẪU tự sinh** để trình diễn quy trình "
        "(tên và giá sản phẩm lấy từ danh mục Bách Hóa Xanh của dự án, còn số bán và tồn kho là mô phỏng). "
        "Không dùng các con số này để ra quyết định thật. Khi có dữ liệu thật từ hệ thống POS, chạy lại "
        "đúng quy trình này là ra báo cáo tương ứng.",
        "",
        f"Phạm vi: {C.N_SKU} sản phẩm, {len(C.STORES)} chi nhánh, dữ liệu theo ngày từ "
        f"{C.START_DATE:%d/%m/%Y} đến {C.END_DATE:%d/%m/%Y} (104 tuần). Dự báo {C.HORIZON_WEEKS} tuần tới "
        f"({k['fc_start']:%d/%m/%Y} – {k['fc_end']:%d/%m/%Y}).",
        "",
        "## 1. Tóm tắt cho quản lý",
        "",
        f"- **Tuần này cần đặt {k['n_order']} mặt hàng, tổng giá trị khoảng {vnd(k['order_value'])}** "
        f"(theo giá vốn). Danh sách chi tiết ở sheet *Đề xuất đặt hàng* của file Excel.",
        f"- **{k['n_urgent']} mặt hàng–chi nhánh có nguy cơ hết hàng**, trong đó {k['n_out']} đang hết hàng "
        f"ngay lúc này. Đây là việc cần xử lý trước tiên.",
        f"- **{len(over)} mặt hàng–chi nhánh tồn kho quá mức** (đủ bán trên {C.OVERSTOCK_WEEKS} tuần), "
        f"đang giữ khoảng {vnd(k['excess_value'])} vốn không cần thiết.",
        f"- Trong 2 năm qua cửa hàng hết hàng khoảng **{pct(k['stockout_rate'])} số ngày bán**, ước tính "
        f"mất khoảng **{vnd(k['lost_revenue'])} doanh thu** vì khách muốn mua mà không có hàng.",
        f"- Mô hình dự báo được chọn sai lệch trung bình **{pct(ho.loc['Chọn theo nhóm', 'WAPE'])}** mỗi tuần "
        f"trên từng mặt hàng, tốt hơn cách đang làm (lấy trung bình 8 tuần gần nhất, sai {pct(ho.loc['MA8', 'WAPE'])}).",
        "",
        "## 2. Phát hiện chính",
        "",
        f"**Xu hướng.** Doanh thu 52 tuần gần nhất đạt {vnd(k['rev_last'])}, "
        f"{'tăng' if k['yoy'] >= 0 else 'giảm'} {pct(abs(k['yoy']))} so với 52 tuần trước đó. "
        f"Có {k['n_growing']} mặt hàng đang tăng trưởng trên 10%/năm và {k['n_declining']} mặt hàng giảm trên 10%/năm.",
        "",
        f"**Mùa vụ.** Tết là đỉnh lớn nhất trong năm: tuần cao điểm trước Tết bán gấp khoảng "
        f"{str(round(k['tet_peak_ratio'], 1)).replace('.', ',')} lần tuần bình thường, rồi giảm mạnh trong tuần nghỉ Tết. Bánh kẹo, bia, "
        f"nước ngọt, dầu ăn, nước mắm tăng mạnh nhất. Nước giải khát bán chạy hơn vào mùa nóng (tháng 3–5), "
        f"mì gói và đồ ăn liền bán tốt hơn vào mùa mưa. Cuối tuần (thứ Bảy, Chủ nhật) bán nhiều hơn ngày thường "
        f"khoảng {pct(k['weekend_uplift'])}. Các đợt 11/11 và 12/12 có khuyến mãi làm số bán tăng rõ rệt.",
        "",
        f"**Hết hàng làm số liệu bán bị thấp hơn nhu cầu thật.** Vào những tuần có hết hàng, số bán ghi nhận "
        f"thấp hơn nhu cầu thật khoảng {pct(abs(k['bias_raw']))}. Sau khi làm sạch (ước lượng lại phần khách "
        f"muốn mua nhưng không có hàng), sai lệch này còn khoảng {pct(abs(k['bias_clean']))}. Nếu dùng thẳng số bán "
        f"để dự báo, cửa hàng sẽ tiếp tục nhập thiếu và lại hết hàng.",
        "",
        f"**Dữ liệu.** Đã xử lý {num(k['missing_record'])} ngày mất bản ghi, {num(k['missing_qty'])} ô số lượng "
        f"bị trống và {num(k['outliers'])} giá trị nhập sai (ví dụ bán nhiều hơn số hàng đang có). Các ngày nghỉ "
        f"Tết được ghi nhận là nghỉ chứ không phải thiếu dữ liệu.",
        "",
        f"**Phân nhóm sản phẩm (ABC-XYZ).** Chỉ {n_a} mặt hàng (nhóm A) đã tạo ra 80% doanh thu. "
        f"Phần lớn mặt hàng có nhu cầu khá ổn định (nhóm X); {k['n_z']} mặt hàng bán rất thất thường hoặc "
        f"lâu lâu mới bán (nhóm Z) – những mặt hàng này khó dự báo, nên giữ tồn kho tối thiểu và đặt theo đơn khi có thể.",
        "",
        "| Nhóm | Ý nghĩa | Số mặt hàng |",
        "|---|---|---|",
    ]
    meaning = {"A": "doanh thu cao", "B": "doanh thu trung bình", "C": "doanh thu thấp",
               "X": "bán đều", "Y": "dao động vừa", "Z": "thất thường"}
    for sname, cnt in seg_counts.sort_index().items():
        lines.append(f"| {sname} | {meaning[sname[0]]}, {meaning[sname[1]]} | {cnt} |")
    lines += [
        "",
        "## 3. Độ chính xác dự báo",
        "",
        "Cách kiểm tra: giả vờ đang ở 3 thời điểm trong quá khứ, chỉ cho mô hình xem dữ liệu trước thời điểm đó, "
        "bắt dự báo 12 tuần tiếp theo rồi so với những gì đã thực sự xảy ra. **WAPE** là tổng sai lệch chia tổng "
        "số bán (WAPE 30% nghĩa là sai trung bình 30 đơn vị trên mỗi 100 đơn vị bán). **Bias** dương nghĩa là dự báo "
        "thường cao hơn thực tế (thừa hàng), âm là thấp hơn (thiếu hàng).",
        "",
        "| Mô hình | Giải thích ngắn | WAPE | Bias |",
        "|---|---|---|---|",
    ]
    desc = {"MA8": "Trung bình 8 tuần gần nhất (cách làm thủ công phổ biến)",
            "SNaive": "Lấy số bán cùng tuần năm trước",
            "ETS": "Làm trơn hàm mũ có xu hướng",
            "TSB": "Croston cải tiến, dành cho hàng bán không đều",
            "LightGBM": "Học máy, hiểu được Tết, khuyến mãi, giá, mùa vụ"}
    for m, r in ov.iterrows():
        lines.append(f"| {m} | {desc[m]} | {pct(r['WAPE'])} | {'+' if r['Bias'] >= 0 else ''}{pct(r['Bias'])} |")
    lines += [
        "",
        f"**{best}** chính xác nhất, giảm sai số khoảng {pct((base_wape - ov.loc[best, 'WAPE']) / base_wape)} so với cách "
        f"lấy trung bình 8 tuần. Cách \"lấy cùng kỳ năm trước\" kém nhất vì Tết mỗi năm rơi vào một tuần khác nhau. "
        f"Mô hình được chọn riêng cho từng nhóm ABC-XYZ (xem sheet *So sánh mô hình*); khi kiểm tra trung thực "
        f"trên giai đoạn cuối (không dùng để chọn mô hình), cách chọn theo nhóm đạt WAPE "
        f"{pct(ho.loc['Chọn theo nhóm', 'WAPE'])}.",
        "",
        f"Độ chính xác khác nhau nhiều giữa các nhóm: nhóm A/X sai khoảng {pct(k['wape_ax'])}, còn nhóm C/Z sai trên "
        f"{pct(k['wape_cz'])} – hàng bán lẻ tẻ thì không mô hình nào đoán chính xác được từng tuần, nên với nhóm này "
        f"phải dựa vào tồn kho an toàn chứ không dựa vào con số dự báo.",
        "",
        "## 4. Khuyến nghị hành động",
        "",
        "**Ngay trong tuần này**",
        "",
        f"1. Đặt hàng theo sheet *Đề xuất đặt hàng* ({k['n_order']} dòng, khoảng {vnd(k['order_value'])}). "
        f"Số lượng đã tính sẵn tồn kho an toàn cho mức phục vụ {C.SERVICE_LEVEL:.0%} và làm tròn theo thùng.",
        "2. Xử lý các mặt hàng nguy cơ hết hàng cao nhất (bảng dưới). Với hàng đang hết, xin nhà cung cấp giao gấp "
        "hoặc điều chuyển từ chi nhánh còn dư.",
        "",
        "| Chi nhánh | Mã | Sản phẩm | Cảnh báo | Tồn | Dự báo tuần tới | Nên đặt |",
        "|---|---|---|---|---|---|---|",
        *[row_u(r) for r in urgent.head(12).itertuples()],
        "",
        "3. Ngừng nhập và cân nhắc khuyến mãi hoặc điều chuyển các mặt hàng tồn dư:",
        "",
        "| Chi nhánh | Mã | Sản phẩm | Số tuần tồn | Vốn tồn dư |",
        "|---|---|---|---|---|",
        *[row_o(r) for r in over.head(10).itertuples()],
        "",
        "**Trong 1–3 tháng tới**",
        "",
        "4. Chạy lại dự báo mỗi tuần (trước buổi đặt hàng thứ Hai) để số đề xuất luôn theo kịp tình hình mới.",
        f"5. Chuẩn bị Tết 2027 (mùng 1 ngày {C.TET_DATES[-1]:%d/%m/%Y}): nhu cầu bắt đầu tăng từ khoảng giữa tháng 1. "
        "Kỳ dự báo này mới phủ đến giữa tháng 1, nên đầu tháng 1 cần chạy lại để có số chính xác cho mùa Tết, "
        "và làm việc sớm với nhà cung cấp bánh kẹo, bia, nước ngọt.",
        f"6. Kiểm tra với nhà cung cấp: {k['n_supplier_issue']} mặt hàng–chi nhánh đang hết hàng mà không có đơn nào "
        "đang trên đường về, cần hỏi lý do (nhà cung cấp hết hàng, giao trễ hay quên đặt).",
        "7. Với nhóm A (ít mặt hàng nhưng mang lại phần lớn doanh thu): ưu tiên không để hết hàng, xem xét đặt "
        "2 lần/tuần. Với nhóm C/Z: giảm tồn, đặt ít và thường xuyên hơn.",
        "8. Ghi nhận đầy đủ ngày hết hàng và kế hoạch khuyến mãi trong POS, vì đây là hai thông tin giúp dự báo "
        "chính xác hơn nhiều.",
        "",
        "## 5. Giả định và giới hạn",
        "",
        *[f"- **{a}:** {b}" for a, b in R["assumptions"]],
        "",
        "## 6. Các file kèm theo",
        "",
        "- `outputs/ket_qua_du_bao_ton_kho.xlsx`: dự báo theo SKU, tồn kho an toàn, ROP, đề xuất đặt hàng, cảnh báo.",
        "- `outputs/charts/top10_thuc_te_vs_du_bao.png`: thực tế và dự báo cho 10 mặt hàng doanh thu cao nhất.",
        "- `outputs/charts/tong_quan_nhu_cau.png`, `outputs/charts/so_sanh_mo_hinh.png`.",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
